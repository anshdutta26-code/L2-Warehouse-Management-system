"""Fill a replaceable Word template from an immutable document snapshot."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import shutil
import subprocess
import tempfile
from docx import Document
from docx.table import _Row

TEMPLATE = Path(__file__).parent / 'templates' / 'invoice.docx'


def _replace(paragraph, values):
    text = paragraph.text
    for key, value in values.items():
        text = text.replace('{{' + key + '}}', str(value))
    if text != paragraph.text:
        if paragraph.runs:
            paragraph.runs[0].text = text
            for run in paragraph.runs[1:]:
                run.text = ''
        else:
            paragraph.add_run(text)


def make_word(d, template=TEMPLATE):
    doc = Document(template)
    money = lambda n: f'{n / 100:,.2f}'
    values = {k: d.get(k, '') for k in ('number', 'date', 'status', 'pos', 'location', 'notes', 'voidReason')}
    values['title'] = 'Tax invoice' if d['kind'] == 'invoice' else 'Unbilled entry'
    values['notice'] = ('Not a tax invoice. Stock issued pending invoicing; tax amounts are estimates.'
                        if d['kind'] == 'unbilled' else '')
    for prefix, source in [('supplier', d['supplier']), ('client', d['client'])]:
        values.update({prefix + '_' + key: source.get(key, '') for key in
                       ('name', 'address', 'city', 'pin', 'state', 'gstin', 'phone', 'bank', 'terms')})
    values.update({key: money(value) for key, value in d['totals'].items()})
    for table in doc.tables:
        for row in list(table.rows):
            if any('{{item}}' in cell.text for cell in row.cells):
                for line in d['lines']:
                    new = deepcopy(row._tr)
                    row._tr.addprevious(new)
                    item = {'item': line['name'], 'hsn': line['hsn'],
                            'qty': f"{line['quantity']:g} {line['unit']}",
                            'rate': money(line['rate']), 'discount': line['discount'],
                            'gst': line['gst'], 'line_taxable': money(line['taxable']),
                            'line_total': money(line['total'])}
                    for cell in _Row(new, table).cells:
                        for paragraph in cell.paragraphs:
                            _replace(paragraph, item)
                table._tbl.remove(row._tr)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace(paragraph, values)
    for paragraph in doc.paragraphs:
        _replace(paragraph, values)
    for section in doc.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                _replace(paragraph, values)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def word_to_pdf(content):
    executable = shutil.which('libreoffice') or shutil.which('soffice')
    if not executable:
        raise RuntimeError('Word preview requires LibreOffice. Install packages.txt on the Streamlit host.')
    with tempfile.TemporaryDirectory(prefix='stockdesk-invoice-') as folder:
        root = Path(folder)
        source = root / 'invoice.docx'
        source.write_bytes(content)
        subprocess.run([executable, '-env:UserInstallation=' + (root / 'profile').as_uri(),
                        '--headless', '--convert-to', 'pdf', '--outdir', str(root), str(source)],
                       check=True, capture_output=True, timeout=60)
        return (root / 'invoice.pdf').read_bytes()


def preview_pages(pdf):
    import fitz
    with fitz.open(stream=pdf, filetype='pdf') as doc:
        return [page.get_pixmap(matrix=fitz.Matrix(1.3, 1.3)).tobytes('png') for page in doc]
